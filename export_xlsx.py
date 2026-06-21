from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models import Project
from naming import generate_labels


HEADERS = ["Oznaczenie", "Port", "Skąd", "Dokąd", "Typ", "Numer", "Status", "Uwagi"]


def export_project_to_xlsx(project: Project, path: str | Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lista kablowa"
    sheet.append(HEADERS)

    for label in generate_labels(project):
        sheet.append(
            [
                label.designation,
                label.port_label,
                label.source,
                label.destination,
                label.export_signal_type,
                label.number,
                label.status,
                label.notes,
            ]
        )

    style_header(sheet)
    autosize_columns(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="2B5C79")
    font = Font(color="EEF8FF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font


def autosize_columns(sheet) -> None:
    for column in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 60)
